#!/usr/bin/env python3
"""Draw a manuscript figure with a central hub and four application panels.

No presentation skill, slide template, title, section number or conclusion bar.
The PPTX is an editable drawing container; PDF/PNG/SVG are the figure outputs.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import shutil
import subprocess

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / 'assets/figure1_v2'
PREVIEWS = ROOT / 'previews'
DOWNLOADS = ROOT / 'downloads'
PREVIEWS.mkdir(exist_ok=True)
DOWNLOADS.mkdir(exist_ok=True)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(12), Inches(11.2)
prs.core_properties.title = '图一（第二版）—电磁场建模在太赫兹领域的应用'
prs.core_properties.subject = '申请书正文配图：中心辐射结构，非演讲幻灯片'
prs.core_properties.author = 'Research Figure'
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')

FONT = 'Noto Serif CJK SC'
INK = '17344C'
BLUE = 'B8DFF4'
OUTLINE = '84B8D7'

def label(text, x, y, w, h, size=24, bold=False):
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    obj.name = text
    tf = obj.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = p.space_after = Pt(0)
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(INK)
    # Set East Asian font explicitly for PowerPoint and LibreOffice.
    from pptx.oxml.xmlchemy import OxmlElement
    ea = OxmlElement('a:ea')
    ea.set('typeface', FONT)
    r._r.get_or_add_rPr().append(ea)
    return obj

def filled_shape(kind, x, y, w, h, fill, line=None):
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        obj.line.color.rgb = RGBColor.from_string(line)
        obj.line.width = Pt(1.25)
    else:
        obj.line.fill.background()
    return obj

def panel(filename, x, y, w, h, title, label_y):
    """Fit the original transparent asset using native picture-crop metadata."""
    src = ASSETS / filename
    with Image.open(src) as im:
        iw, ih = im.size
        if im.mode != 'RGBA':
            raise RuntimeError(f'Expected reviewed RGBA asset: {src}')
        alpha = im.getchannel('A')
        bbox = alpha.point(lambda a: 255 if a > 12 else 0).getbbox()
        if not bbox:
            raise RuntimeError(f'Empty asset: {src}')
    left, top, right, bottom = bbox
    cw, ch = right-left, bottom-top
    scale = min(w/cw, h/ch)
    pw, ph = cw*scale, ch*scale
    pic = slide.shapes.add_picture(str(src), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2),
                                   width=Inches(pw), height=Inches(ph))
    pic.crop_left, pic.crop_right = left/iw, (iw-right)/iw
    pic.crop_top, pic.crop_bottom = top/ih, (ih-bottom)/ih
    pic.name = title + '（概念示意子图）'
    pic._element.nvPicPr.cNvPr.set('descr', title+'科研概念示意，非测量或仿真结果。')
    label(title, x, label_y, w, 0.46)

# Four application subfigures, in the cardinal directions of the user's example.
panel('communication.png', 3.90, 0.20, 4.20, 2.80, '高速通信', 3.04)
panel('security.png', 0.12, 4.05, 3.82, 2.55, '安检成像', 6.64)
panel('biosensing.png', 8.06, 4.05, 3.82, 2.55, '生物传感', 6.64)
panel('material_identification.png', 3.90, 7.92, 4.20, 2.80, '物质识别', 10.69)

# Short, equal-weight arrows lead out from the central modelling hub.
arrows = [
    (MSO_AUTO_SHAPE_TYPE.UP_ARROW, 5.69, 3.57, 0.62, 0.79),
    (MSO_AUTO_SHAPE_TYPE.LEFT_ARROW, 4.00, 5.29, 0.71, 0.62),
    (MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, 7.29, 5.29, 0.71, 0.62),
    (MSO_AUTO_SHAPE_TYPE.DOWN_ARROW, 5.69, 6.84, 0.62, 0.84),
]
for i, (kind, x, y, w, h) in enumerate(arrows, 1):
    obj = filled_shape(kind, x, y, w, h, BLUE)
    obj.name = f'中心向外箭头{i}'
    # A broad arrowhead and a compact shaft, matching the reference structure.
    obj.adjustments[0] = 0.52
    obj.adjustments[1] = 0.42

hub = filled_shape(MSO_AUTO_SHAPE_TYPE.OVAL, 4.65, 4.25, 2.70, 2.70, 'D8EDF9', OUTLINE)
hub.name = '中心：电磁场建模'
label('电磁场建模', 4.71, 5.23, 2.58, 0.66, size=29, bold=True)

slide.notes_slide.notes_text_frame.text = (
    '图1（第二版）：电磁场建模在太赫兹领域中的应用。正文配图，采用用户参考图的中心圆与四向辐射结构。\n'
    '四周应用：高速通信、安检成像、生物传感、物质识别，均来自2026年10月22日最新申请书“需求分析”首段。\n'
    '本图不使用演讲幻灯片的标题、分栏编号、结论条或PPT制作skill。\n'
    '四张独立子图均由OpenAI图像工具生成，为概念示意，不是实验照片、测量结果或仿真结果。'
    '波纹、热点、分子/滴液尺度、物质识别谱线和穿透成像效果仅定性表达应用场景，不构成定量性能证据。\n'
    '中心圆、箭头和五处中文标签均为原生可编辑元素。未改动第一版或原始申请书。'
)
out = ROOT / '图一_第二版.pptx'
prs.save(out)

subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(PREVIEWS), str(out)],
               cwd=ROOT, check=True)
pdf_file = PREVIEWS / '图一_第二版.pdf'
with fitz.open(pdf_file) as pdf:
    assert len(pdf) == 1
    page = pdf[0]
    content = page.get_text()
    for t in ['电磁场建模', '高速通信', '安检成像', '生物传感', '物质识别']:
        assert t in content, f'Missing rendered label: {t}'
    page.get_pixmap(dpi=300, alpha=False).save(PREVIEWS / '图一_第二版.png')
    page.get_pixmap(dpi=110, alpha=False).save(PREVIEWS / '图一_第二版_预览.png')
    (PREVIEWS / '图一_第二版.svg').write_text(page.get_svg_image(text_as_path=True))

for source, name in [
    (out, 'figure1_v2.pptx'),
    (pdf_file, 'figure1_v2.pdf'),
    (PREVIEWS / '图一_第二版.png', 'figure1_v2.png'),
    (PREVIEWS / '图一_第二版_预览.png', 'figure1_v2_preview.png'),
    (PREVIEWS / '图一_第二版.svg', 'figure1_v2.svg'),
]:
    target = DOWNLOADS / name
    shutil.copyfile(source, target)
    target.chmod(0o644)

with ZipFile(DOWNLOADS / 'figure1_v2.zip', 'w', ZIP_DEFLATED) as z:
    for name in ['figure1_v2.pptx', 'figure1_v2.pdf', 'figure1_v2.png', 'figure1_v2.svg']:
        z.write(DOWNLOADS / name, name)
(DOWNLOADS / 'figure1_v2.zip').chmod(0o644)

print(json.dumps({
    'figure': str(out), 'pictures': 4, 'native_labels': 5,
    'size_inches': [12, 11.2], 'png_pixels': [3600, 3360],
    'outputs': [p.name for p in DOWNLOADS.glob('figure1_v2*')],
}, ensure_ascii=False, indent=2))
