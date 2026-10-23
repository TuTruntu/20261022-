#!/usr/bin/env python3
"""A six-branch manuscript figure, drawn independently of presentation skills."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import argparse
import math
import shutil
import subprocess

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as ST
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
PREVIEW = ROOT / 'previews'
OUT = ROOT / 'downloads'
PREVIEW.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)
args = argparse.ArgumentParser()
args.add_argument('--preview-only', action='store_true')
args = args.parse_args()

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(12.95)
prs.core_properties.title = '图一第三版：太赫兹电磁建模的应用与作用'
prs.core_properties.subject = '正文配图；六向应用辐射与器件—场—响应关系'
prs.core_properties.author = 'Research Figure'
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
INK, BLUE, PALE, GOLD, GRAY = '14394F', '2589B3', 'B8DEEF', 'D79B3A', '536F7F'
FONT = 'Noto Serif CJK SC'

def fill(obj, color, stroke=None, width=1):
    obj.fill.solid()
    obj.fill.fore_color.rgb = RGBColor.from_string(color)
    if stroke:
        obj.line.color.rgb = RGBColor.from_string(stroke)
        obj.line.width = Pt(width)
    else:
        obj.line.fill.background()
    # Suppress theme-inherited shadows for a clean manuscript figure.
    obj._element.spPr.append(OxmlElement('a:effectLst'))
    return obj

def shape(kind, x, y, w, h, color, stroke=None, width=1):
    return fill(s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)), color, stroke, width)

def poly(points, color, stroke=None, width=1, name='原生矢量科研示意'):
    builder = s.shapes.build_freeform(Inches(points[0][0]), Inches(points[0][1]))
    builder.add_line_segments([(Inches(x), Inches(y)) for x, y in points[1:]], close=True)
    obj = builder.convert_to_shape()
    obj.name = name
    return fill(obj, color, stroke, width)

def line(x1, y1, x2, y2, color=BLUE, width=1):
    obj = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb = RGBColor.from_string(color)
    obj.line.width = Pt(width)
    return obj

def txt(text, x, y, w, h, size=24, color=INK, bold=False, align=PP_ALIGN.CENTER):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    obj.name = text
    tf = obj.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name, run.font.size = FONT, Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    ea = OxmlElement('a:ea')
    ea.set('typeface', FONT)
    run._r.get_or_add_rPr().append(ea)
    return obj

def arrow(p, q, color=PALE, shaft=0.11, head=0.24, head_width=0.31):
    dx, dy = q[0]-p[0], q[1]-p[1]
    norm = math.hypot(dx,dy)
    ux, uy = dx/norm, dy/norm
    vx, vy = -uy, ux
    b = (q[0]-head*ux, q[1]-head*uy)
    add=lambda r,a:(r[0]+a*vx,r[1]+a*vy)
    return poly([add(p,shaft/2),add(b,shaft/2),add(b,head_width/2),q,
                 add(b,-head_width/2),add(b,-shaft/2),add(p,-shaft/2)],color)

def picture(src, x, y, w, h, description):
    with Image.open(src) as im:
        iw, ih = im.size
        bbox = im.getchannel('A').point(lambda a:255 if a>12 else 0).getbbox() if im.mode=='RGBA' else (0,0,iw,ih)
    l,t,r,b=bbox
    cw,ch=r-l,b-t
    sc=min(w/cw,h/ch)
    pw,ph=cw*sc,ch*sc
    pic=s.shapes.add_picture(str(src),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),Inches(pw),Inches(ph))
    pic.crop_left,pic.crop_right=l/iw,(iw-r)/iw
    pic.crop_top,pic.crop_bottom=t/ih,(ih-b)/ih
    pic.name=description
    pic._element.nvPicPr.cNvPr.set('descr',description+'；概念示意，非测量或仿真结果。')
    return pic

def application(src,x,y,w,h,title,subtitle,title_y):
    picture(src,x,y,w,h,title)
    tx,sx=max(0,x-0.15),max(0,x-0.25)
    txt(title,tx,title_y,min(w+0.30,16-tx),0.49,26,bold=True)
    txt(subtitle,sx,title_y+0.51,min(w+0.50,16-sx),0.36,19,GRAY)

def field_color(v):
    stops=[(0,(39,96,165)),(.35,(44,177,207)),(.65,(144,218,190)),(.84,(250,205,94)),(1,(233,114,50))]
    v=max(0,min(1,v))
    for (a,ca),(b,cb) in zip(stops,stops[1:]):
        if a<=v<=b:
            t=(v-a)/(b-a)
            return ''.join(f'{round(x+(y-x)*t):02X}' for x,y in zip(ca,cb))
    return 'E97232'

def field_cube(x,y,w,h):
    # Three qualitative cross-sections; no axes values or empirical data.
    def P(a,b,c):return (x+w*(.06+.69*a+.25*b),y+h*(.94-.22*b-.67*c))
    for axis,fixed in [('z',.08),('y',.55),('x',.55)]:
        n=7
        for i in range(n):
            for j in range(n):
                u0,u1=i/n,(i+1)/n
                v0,v1=j/n,(j+1)/n
                def point(u,v):
                    return P(u,v,fixed) if axis=='z' else P(u,fixed,v) if axis=='y' else P(fixed,u,v)
                u,v=(u0+u1)/2,(v0+v1)/2
                value=.18+.73*math.exp(-((u-.55)**2/.030+(v-.46)**2/.13))+.13*math.sin(3*math.pi*u)**2
                poly([point(u0,v0),point(u1,v0),point(u1,v1),point(u0,v1)],field_color(value))
    corners=[P(a,b,c) for a,b,c in [(0,0,0),(1,0,0),(1,1,0),(0,1,0),(0,0,1),(1,0,1),(1,1,1),(0,1,1)]]
    for i,j in [(0,1),(1,2),(2,3),(3,0),(4,5),(5,6),(6,7),(7,4),(0,4),(1,5),(2,6),(3,7)]:
        line(*corners[i],*corners[j],'7FAABD',.7)
    for xx,yy in [(x+.40*w,y+.54*h),(x+.63*w,y+.70*h),(x+.78*w,y+.42*h)]:
        arrow((xx,yy),(xx+.13*w,yy-.10*h),'FFFFFF',shaft=.014,head=.045,head_width=.05)

def response_plot(x,y,w,h):
    # Native spectrum and radiation-pattern symbols; deliberately no numbers.
    line(x+.08*w,y+.70*h,x+.97*w,y+.70*h,'A5B9C5',.8)
    line(x+.08*w,y+.08*h,x+.08*w,y+.70*h,'A5B9C5',.8)
    pts=[]
    for i in range(65):
        u=i/64
        amp=.18+.48*math.exp(-((u-.38)/.055)**2)+.34*math.exp(-((u-.72)/.075)**2)
        pts.append((x+(.08+.89*u)*w,y+amp*h))
    for p,q in zip(pts,pts[1:]):line(*p,*q,BLUE,1.25)
    cx,cy=x+.67*w,y+.93*h
    for scale,col in [(1,GOLD),(.72,BLUE)]:
        pts=[]
        for i in range(61):
            t=2*math.pi*i/60
            r=(.55+.45*math.cos(2*t)**2)*.24*w*scale
            pts.append((cx+r*math.cos(t),cy+.56*r*math.sin(t)))
        for p,q in zip(pts,pts[1:]):line(*p,*q,col,1.15)

# Six directions: the template's radial organisation, expanded beyond four spokes.
for a,b in [((6.10,4.55),(5.30,3.98)),((9.90,4.55),(10.70,3.98)),
            ((4.90,6.40),(4.30,6.40)),((11.10,6.40),(11.70,6.40)),
            ((6.12,8.35),(5.30,9.13)),((9.88,8.35),(10.70,9.13))]:
    arrow(a,b,shaft=.15,head=.26,head_width=.36)

# The hub contains a scientific relationship, rather than an empty labelled disc.
hub=shape(ST.OVAL,4.85,4.14,6.30,4.65,'FFFFFF','99C7DF',1.4)
hub.name='电磁建模核心：结构—场—响应'
txt('太赫兹电磁场建模',5.30,4.50,5.40,.57,30,bold=True)
txt('幅度 · 相位 · 极化 · 局域场',5.15,5.08,5.70,.34,19,GRAY)

for t,x in [('器件结构',5.10),('三维电磁场',7.15),('电磁响应',9.20)]:
    txt(t,x,5.61,1.70,.40,21.5,bold=True)
picture(ROOT/'assets/figure1_v3/core_device.png',5.25,6.04,1.67,1.35,'亚波长开口环阵列')
field_cube(7.23,6.10,1.56,1.30)
response_plot(9.27,6.05,1.55,1.30)
arrow((6.88,6.73),(7.17,6.73),'58A3C2',shaft=.035,head=.10,head_width=.14)
arrow((8.86,6.73),(9.15,6.73),'58A3C2',shaft=.035,head=.10,head_width=.14)
txt('几何·材料',5.22,7.35,1.98,.31,17.5,GRAY)
txt('激励条件',5.22,7.68,1.98,.28,16.5,GRAY)
txt('E / H · 局域场',7.03,7.37,1.95,.31,17.5,GRAY)
txt('光谱·方向图',9.07,7.37,1.95,.31,17.5,GRAY)

# Response-to-structure return path: modelling supports repeated design decisions.
line(10.16,7.98,10.16,8.17,'7FAEC5',1.1)
line(10.16,8.17,5.83,8.17,'7FAEC5',1.1)
arrow((5.83,8.17),(5.83,7.98),'7FAEC5',shaft=.025,head=.07,head_width=.10)
shape(ST.RECTANGLE,7.23,8.00,1.54,.34,'FFFFFF')
txt('设计优化',7.24,8.00,1.52,.34,18.5,GRAY)

A2=ROOT/'assets/figure1_v2'
A3=ROOT/'assets/figure1_v3'
application(A2/'communication.png',1.12,.18,4.65,2.78,'高速通信','波束与极化调控',3.06)
application(A2/'security.png',10.23,.18,4.65,2.78,'安检成像','穿透成像与对比增强',3.06)
application(A3/'situational_sensing.png',.12,4.79,4.10,2.70,'态势感知','散射分析与成像',7.57)
application(A2/'biosensing.png',11.78,4.79,4.10,2.70,'生物传感','局域场增强与感知',7.57)
application(A3/'nondestructive_testing.png',1.12,9.21,4.65,2.58,'无损检测','缺陷与界面探测',11.91)
application(A2/'material_identification.png',10.23,9.21,4.65,2.58,'物质识别','光谱指纹识别',11.91)

s.notes_slide.notes_text_frame.text=(
    '正文图一第三版。用户要求保留中心辐射组织但增加内容，不限制上下左右四个方向。\n'
    '外圈六类场景均来自最新申请书需求分析首段：高速通信、安检成像、生物传感、物质识别、无损检测、态势感知。\n'
    '中心关系：在给定几何、材料和激励条件下求取三维电磁场，再提取光谱、方向图等电磁响应，支持器件设计优化。'
    '幅度、相位、极化和局域场是文中明确的器件调控维度。\n'
    '全部场景、器件和谱线均为概念示意；中心三维场切片是定性原生矢量图，不是求解器输出。'
    '外圈六张图和中心器件图片由OpenAI图像工具生成；材料检测使用介质层压复合材料，而非宣称太赫兹可穿透金属。'
    '态势感知只表示散射与成像潜力，不作远程、全天候或既成系统能力承诺。\n'
    '本图不使用PPT制作skill，不包含演讲标题、章节编号或结论条；中心、文字、连接、场切片和响应曲线可编辑。'
    '项目拟考核性能指标没有被表述为已实现结果。第一版、第二版和原始Word均保持不变。'
)

# Remove theme effect references as well as explicit shadows.
for obj in s.shapes:
    for child in list(obj._element):
        if child.tag == qn('p:style'):
            obj._element.remove(child)

target=ROOT/'图一_第三版.pptx'
prs.save(target)
subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(PREVIEW),str(target)],check=True,cwd=ROOT)
with fitz.open(PREVIEW/'图一_第三版.pdf') as pdf:
    assert len(pdf)==1
    page=pdf[0]
    text=page.get_text()
    for expected in ['高速通信','安检成像','生物传感','物质识别','无损检测','态势感知','电磁场建模','设计优化']:
        assert expected in text,expected
    page.get_pixmap(dpi=105,alpha=False).save(PREVIEW/'图一_第三版_预览.png')
    if not args.preview_only:
        page.get_pixmap(dpi=300,alpha=False).save(PREVIEW/'图一_第三版.png')
        (PREVIEW/'图一_第三版.svg').write_text(page.get_svg_image(text_as_path=True))
if args.preview_only:
    print('Preview:',PREVIEW/'图一_第三版_预览.png')
else:
    for src,name in [(target,'figure1_v3.pptx'),(PREVIEW/'图一_第三版.pdf','figure1_v3.pdf'),
                     (PREVIEW/'图一_第三版.png','figure1_v3.png'),(PREVIEW/'图一_第三版_预览.png','figure1_v3_preview.png'),
                     (PREVIEW/'图一_第三版.svg','figure1_v3.svg')]:
        dst=OUT/name
        shutil.copyfile(src,dst)
        dst.chmod(0o644)
    with ZipFile(OUT/'figure1_v3.zip','w',ZIP_DEFLATED) as z:
        for name in ['figure1_v3.pptx','figure1_v3.pdf','figure1_v3.png','figure1_v3.svg']:
            z.write(OUT/name,name)
    print('Exported manuscript figure v3:',OUT/'figure1_v3.png')
