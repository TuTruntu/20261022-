#!/usr/bin/env python3
"""Check five-slide scope, editable text, page geometry and actual PDF text."""
from pathlib import Path
from zipfile import ZipFile
import re
import hashlib
import fitz
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'downloads/presentation_5slides'
PPTX=OUT/'thz_background_objectives_5slides_revised.pptx'
PDF=PPTX.with_suffix('.pdf')
pres=Presentation(PPTX)
assert len(pres.slides)==5
titles=['1.1 研究背景及现状','1.2 现有电磁建模方法','1.3 现有研究面临的具体挑战',
        '2.1 研究目标','2.2 研究目标实现面临的困难']
norm=lambda s:re.sub(r'\s+','',s)
counts=[]
page_text=[]
with fitz.open(PDF) as pdf:
    assert len(pdf)==5
    for i,slide in enumerate(pres.slides):
        labels=[s.text for s in slide.shapes if s.has_text_frame and s.text.strip()]
        assert titles[i] in labels
        assert len(labels)>15, 'Slides must have native, editable text.'
        assert slide.has_notes_slide
        content=norm(pdf[i].get_text())
        page_text.append(content)
        missing=[t for t in labels if norm(t) not in content]
        assert not missing,(i+1,missing)
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0,(i+1,shape.name)
            assert shape.left+shape.width<=pres.slide_width+100,(i+1,shape.name)
            assert shape.top+shape.height<=pres.slide_height+100,(i+1,shape.name)
        for word in pdf[i].get_text('words'):
            x0,y0,x1,y1,*_=word
            assert x0>=-.1 and y0>=-.1 and x1<=pdf[i].rect.width+.1 and y1<=pdf[i].rect.height+.1,(i+1,word)
        counts.append(len(labels))
        # The five-page opening must not assert the literature's numerical results.
        assert not any(x in content for x in ['88%','75%','100倍','26%','≤5%','≥20种'])
with ZipFile(PPTX) as z:
    assert z.testzip() is None
    assert len([n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)])==5
    with ZipFile(OUT/'thz_background_objectives_5slides.pptx') as previous:
        for number in [4,5]:
            part=f'ppt/slides/slide{number}.xml'
            assert z.read(part)==previous.read(part),f'Slide {number} was meant to stay unchanged.'
# Guard the revised narrative separation requested by the user.
for phrase in ['全波仿真','响应建模','神经算子','FNO','FDTD','弊端','局限','瓶颈']:
    assert phrase not in page_text[0],('Background must not classify or critique methods',phrase)
for phrase in ['优势','弊端','局限','不足','受限','退化','瓶颈','挑战','缺失']:
    assert phrase not in page_text[1],('Method page must remain descriptive',phrase)
for phrase in ['全波仿真','电磁响应建模','FDTD','FEM','MoM','DeepONet','FNO','覆盖率','距离场']:
    assert phrase not in page_text[2],('Challenge page must not repeat the method overview',phrase)
source=ROOT/'docs/presentation_5slides/太赫兹国重申请书-汇总-10-22.docx'
assert hashlib.sha256(source.read_bytes()).hexdigest()=='15d8ce5a779f0d1df8fd64349a37fbcf44769be8e36d077b24ab54c103c6b70c'
print('PASS: five pages; native text/PDF and bounds; background/methods/challenges separated; slides 4 and 5 exactly unchanged; sources and DOCX intact.')
print('Native text objects per page:',counts)
